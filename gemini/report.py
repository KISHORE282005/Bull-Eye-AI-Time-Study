import json
import re
import pandas as pd
from pathlib import Path

from utils.export import write_overall_analysis_sheet, to_operation_text

# =========================================================
# OUTPUT FOLDER
# =========================================================

# Anchored to the project root, never to the current working
# directory, so the reports always land in the same /output the
# dashboard reads from - whatever folder Streamlit was started in.

OUTPUT = Path(__file__).resolve().parent.parent / "output"
OUTPUT.mkdir(parents=True, exist_ok=True)

JSON_FILE = OUTPUT / "time_study.json"
CSV_FILE = OUTPUT / "activities.csv"

# =========================================================
# SANITIZE VIDEO NAME
# =========================================================

def sanitize_video_name(video_name):
    """
    Convert a video file name into a safe report base name.
    "Site_A_Crane_Lift.MTS" -> "Site_A_Crane_Lift"
    """

    base = Path(video_name).stem if video_name else "Industrial_Time_Study"

    base = re.sub(r'[<>:"/\\|?*]', "_", str(base))

    base = base.strip().rstrip(".")

    return base or "Industrial_Time_Study"

# =========================================================
# SAVE REPORT
# =========================================================

def save_report(data, video_name=None):
    """
    Save JSON / CSV / Excel report.

    The Excel workbook is named after the source video, e.g.
    output/Site_A_Crane_Lift_Time_Study_Report.xlsx
    """

    # -------------------------------------------------
    # Excel file name based on the video name
    # -------------------------------------------------

    base = sanitize_video_name(video_name)

    excel_file = OUTPUT / f"{base}_Time_Study_Report.xlsx"

    data["video_file_name"] = video_name or ""

    data["excel_report_file"] = excel_file.name

    # -------------------------------------------------
    # Save JSON
    # -------------------------------------------------

    with open(JSON_FILE, "w", encoding="utf-8") as f:

        json.dump(
            data,
            f,
            indent=4,
            ensure_ascii=False
        )

    # -------------------------------------------------
    # Activities DataFrame
    # -------------------------------------------------

    activities = pd.DataFrame(
        data.get("activities", [])
    )

    # -------------------------------------------------
    # Required Columns
    # -------------------------------------------------

    required_columns = [

        "process_no",

        "process_name",

        "process_operation",

        "process_description",

        "start_timestamp",

        "end_timestamp",

        "duration",

        "op1",
        "op2",
        "op3",
        "op4",
        "op5",

        "op_wt1",
        "op_wt2",
        "op_wt3",
        "op_wt4",
        "op_wt5",

        "toct",

        "va",

        "nva",

        "r_nva",

        "nva_category",

        "nva_reason"

    ]

    # -------------------------------------------------
    # Add Missing Columns
    # -------------------------------------------------

    numeric_columns = [

        "duration",

        "op1",
        "op2",
        "op3",
        "op4",
        "op5",

        "op_wt1",
        "op_wt2",
        "op_wt3",
        "op_wt4",
        "op_wt5",

        "toct",
        "va",
        "nva",
        "r_nva"

    ]

    for column in required_columns:

        if column not in activities.columns:

            if column in numeric_columns:

                activities[column] = 0.0

            else:

                activities[column] = ""

    # -------------------------------------------------
    # Arrange Columns
    # -------------------------------------------------

    activities = activities[
        required_columns
    ].copy()

    activities["process_description"] = (
        activities["process_description"].map(to_operation_text)
    )

    # -------------------------------------------------
    # Rename Excel Columns
    # -------------------------------------------------

    activities.columns = [

        "Process No",

        "Process Name",

        "Process Operation",

        "Process Description",

        "Start Timestamp",

        "End Timestamp",

        "Duration (sec)",

        "Op1 (sec)",
        "Op2 (sec)",
        "Op3 (sec)",
        "Op4 (sec)",
        "Op5 (sec)",

        "WT1 (sec)",
        "WT2 (sec)",
        "WT3 (sec)",
        "WT4 (sec)",
        "WT5 (sec)",

        "TOCT (sec)",

        "VA (sec)",

        "NVA (sec)",

        "Required NVA (sec)",

        "NVA Category",

        "NVA Reason"

    ]
    # -------------------------------------------------
    # Save Excel
    # -------------------------------------------------

    with pd.ExcelWriter(
        excel_file,
        engine="openpyxl"
    ) as writer:

        # ---------------------------------------------
        # Time Study Sheet
        # ---------------------------------------------

        activities.to_excel(
            writer,
            sheet_name="Time Study",
            index=False
        )

        workbook = writer.book
        worksheet = writer.sheets["Time Study"]

        from openpyxl.styles import Font, PatternFill, Alignment

        # ---------------------------------------------
        # Header Style
        # ---------------------------------------------

        header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="C00000"
        )

        center = Alignment(
            horizontal="center",
            vertical="center"
        )

        for cell in worksheet[1]:

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        # ---------------------------------------------
        # Auto Fit Columns
        # ---------------------------------------------

        for column_cells in worksheet.columns:

            length = max(
                len(str(cell.value))
                if cell.value else 0
                for cell in column_cells
            )

            worksheet.column_dimensions[
                column_cells[0].column_letter
            ].width = length + 5

        # ---------------------------------------------
        # Freeze Header
        # ---------------------------------------------

        worksheet.freeze_panes = "A2"

        # ---------------------------------------------
        # Wrap the description so the sheet stays readable
        # ---------------------------------------------

        wrap = Alignment(
            horizontal="left",
            vertical="top",
            wrap_text=True
        )

        description_column = (
            list(activities.columns).index("Process Description") + 1
        )

        for row in range(2, worksheet.max_row + 1):

            worksheet.cell(
                row=row,
                column=description_column
            ).alignment = wrap

        worksheet.column_dimensions[
            worksheet.cell(row=1, column=description_column).column_letter
        ].width = 80

        # ---------------------------------------------
        # Filter buttons so NVA rows can be picked out
        # on the Time Study sheet as well
        # ---------------------------------------------

        worksheet.auto_filter.ref = worksheet.dimensions

        # ---------------------------------------------
        # Overall Analysis Sheet
        #
        # Total Time / VA / NVA, the NVA breakdown by
        # condition, and the list of exactly which
        # activities the NVA is made of.
        # ---------------------------------------------

        overall_sheet = workbook.create_sheet("Overall Analysis")

        # Process number -> its row on the Time Study sheet,
        # so every NVA activity links back to its own row.

        time_study_rows = {
            process_no: index + 2
            for index, process_no in enumerate(
                activities["Process No"].tolist()
            )
        }

        write_overall_analysis_sheet(

            overall_sheet,

            data.get("overall_analysis", {}),

            data.get("nva_breakdown", {}),

            total_processes=data.get(
                "total_processes",
                len(activities)
            ),

            time_study_sheet="Time Study",

            time_study_rows=time_study_rows

        )

    # -------------------------------------------------
    # Save CSV
    # -------------------------------------------------

    activities.to_csv(

        CSV_FILE,

        index=False

    )

    # -------------------------------------------------
    # Console Messages
    # -------------------------------------------------

    print("=" * 60)

    print("Industrial AI Time Study Report Generated")

    print("=" * 60)

    print(f"JSON      : {JSON_FILE}")

    print(f"Excel     : {excel_file}")

    print(f"CSV       : {CSV_FILE}")

    print(f"Detected  : {data.get('operator_count', 1)} operator(s)")

    print("=" * 60)

    return activities
